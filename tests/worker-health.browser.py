import asyncio, json, shutil, tempfile, os, threading
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from functools import partial
from pathlib import Path
from zipfile import ZipFile
import openpyxl
from PIL import Image
from playwright.async_api import async_playwright, expect

work=Path(tempfile.mkdtemp(prefix='worker-health-tests-'))
ext=Path(os.environ['EDGE_EXTENSION_DIR']) if os.environ.get('EDGE_EXTENSION_DIR') else None
chromium=os.environ.get('CHROMIUM', '/usr/bin/chromium')
base_url=os.environ.get('BASE_URL', 'http://127.0.0.1:3004')
def fixture(path, count=40, pressure=True):
 wb=openpyxl.Workbook();ws=wb.active;ws.title='근로자'
 ws.append(['열'+str(i) for i in range(17 if pressure else 12)])
 data=[('고령'+str(i+1).zfill(2),'631024-1','180/110') for i in range(count)]
 data += [('수축기경계','631025-1','150/99'),('이완기경계','000101-3','149/100'),('정상','000101-3','149/99'),('나이미상','','180/110'),('<img src=x onerror=alert(1)>','000101-3','2026/10/16')]
 for name,birth,bp in data:
  row=['']*(17 if pressure else 12);row[3]='대한건설';row[4]='전기';row[5]=name;row[9]='한국';row[11]=birth
  if pressure:row[16]=bp
  ws.append(row)
 wb.create_sheet('빈 시트');wb.save(path)
fixture(work/'workers.xlsx');fixture(work/'no-q.xlsx',1,False)
(work/'bad.xlsx').write_bytes(b'not an excel file')

async def flow(page, label):
 errors=[];remote=[]
 page.on('pageerror', lambda e:errors.append(str(e)))
 page.on('request', lambda r:remote.append(r.url) if r.url.startswith('https://') else None)
 await page.locator('[data-wh=file]').wait_for()
 await page.locator('[data-wh=asof]').fill('2026-10-24');await page.locator('[data-wh=asof]').dispatch_event('change')
 await page.locator('[data-wh=start]').fill('2026-10-26');await page.locator('[data-wh=start]').dispatch_event('change')
 await page.locator('[data-wh=end]').fill('2026-10-27');await page.locator('[data-wh=end]').dispatch_event('change')
 await page.locator('[data-wh=file]').set_input_files(work/'workers.xlsx')
 await page.locator('[data-wh=metrics]').filter(has_text='45명').wait_for()
 assert await page.locator('[data-wh=body] tr').count()==40
 assert '40명' in await page.locator('[data-wh=metrics]').inner_text()
 assert '1명' in await page.locator('[data-wh=warnings]').inner_text()
 async with page.expect_download() as info:await page.locator('[data-wh=excel]').click()
 download=await info.value;await download.save_as(work/f'{label}.xlsx')
 wb=openpyxl.load_workbook(work/f'{label}.xlsx');ws=wb.active
 assert ws['A1'].value=='2026.10.26~10.27 고령자 혈압측정'
 assert ws.max_row==52
 assert ws['C3'].value=='고령01'
 assert str(ws.print_title_rows)=='$1:$2'
 assert ws['E3'].fill.fgColor.rgb=='FFFCE4D6'
 await page.locator('[data-wh=images]').click();await page.wait_for_selector('[data-wh=previews] a',timeout=60000)
 assert await page.locator('[data-wh=previews] img').count()==2
 await page.locator('[data-wh=previews] img').first.scroll_into_view_if_needed()
 await expect(page.locator('[data-wh=previews] img').first).to_have_js_property('naturalWidth',2480)
 async with page.expect_download() as info:await page.locator('[data-wh=previews] a').first.click()
 download=await info.value;await download.save_as(work/f'{label}.png')
 assert Image.open(work/f'{label}.png').size==(2480,3508)
 await page.locator('[data-wh=pressure]').click()
 assert await page.locator('[data-wh=body] tr').count()==2
 names=await page.locator('[data-wh=body]').inner_text()
 assert '수축기경계' in names and '이완기경계' in names and '고령' not in names
 assert await page.locator('[data-wh=previews] img').count()==0
 await page.locator('[data-wh=end]').fill('2026-10-25');await page.locator('[data-wh=end]').dispatch_event('change')
 assert await page.locator('[data-wh=excel]').is_disabled()
 assert '종료일' in await page.locator('[data-wh=error]').inner_text()
 await page.locator('[data-wh=end]').fill('2026-10-27');await page.locator('[data-wh=end]').dispatch_event('change')
 assert not await page.locator('[data-wh=excel]').is_disabled()
 await page.locator('[data-wh=sheet]').select_option('빈 시트')
 assert await page.locator('[data-wh=result]').is_hidden()
 assert '비어' in await page.locator('[data-wh=error]').inner_text()
 await page.locator('[data-wh=sheet]').select_option('근로자')
 assert await page.locator('[data-wh=body] tr').count()==2
 await page.locator('[data-wh=file]').set_input_files(work/'no-q.xlsx')
 await page.locator('[data-wh=warnings]').filter(has_text='Q열').wait_for()
 assert await page.locator('[data-wh=excel]').is_disabled()
 await page.locator('[data-wh=older]').click();assert not await page.locator('[data-wh=excel]').is_disabled()
 await page.locator('[data-wh=file]').set_input_files(work/'bad.xlsx')
 await expect(page.locator('[data-wh=error]')).not_to_be_empty()
 assert await page.locator('[data-wh=result]').is_hidden()
 await page.locator('[data-wh=file]').set_input_files(work/'workers.xlsx')
 await page.wait_for_selector('[data-wh=result]:not([hidden])')
 await page.set_viewport_size({'width':390,'height':844})
 assert await page.evaluate('document.documentElement.scrollWidth<=innerWidth'),label+' mobile overflow'
 await page.screenshot(path=str(work/f'{label}-mobile.png'),full_page=True)
 await page.locator('[data-wh=clear]').click();assert await page.locator('[data-wh=result]').is_hidden()
 assert not errors,errors
 assert not remote,remote
 print(label,'PASS upload → classification → styled XLSX → 2 PNG pages, bad date/file/sheet, missing Q, reset, 390px; no remote requests')

async def main():
 async with async_playwright() as p:
  browser=await p.chromium.launch(executable_path=chromium,args=['--no-sandbox'])
  page=await browser.new_page(viewport={'width':1440,'height':1000})
  await page.goto(base_url+'/?preview=worker-health')
  await flow(page,'web')
  await page.evaluate("showPage('bp')");assert await page.locator('#bpMonday').is_visible()
  await browser.close()
  if not ext:
   print('Extension test skipped: set EDGE_EXTENSION_DIR to enable it');return
  # Execute the exact packaged extension page with the manifest's CSP.
  # This checks the distributed bundle but does not claim an Edge installation test.
  class Handler(SimpleHTTPRequestHandler):
   def end_headers(self):
    self.send_header('Content-Security-Policy', "script-src 'self'; object-src 'none'")
    super().end_headers()
   def log_message(self,*args): pass
  server=ThreadingHTTPServer(('127.0.0.1',0),partial(Handler,directory=str(ext)))
  threading.Thread(target=server.serve_forever,daemon=True).start()
  browser=await p.chromium.launch(executable_path=chromium,args=['--no-sandbox'])
  page=await browser.new_page()
  await page.add_init_script("""window.chrome = { storage: { local: { get: async () => ({}), set: async () => {} } }, runtime: { getURL: path => path }, tabs: { create: args => { window.openedTool = args.url; } } };""")
  base='http://127.0.0.1:'+str(server.server_port)
  await page.goto(base+'/popup.html');await page.locator('#workerHealth').click()
  assert await page.evaluate('window.openedTool')=='worker-health.html'
  await page.goto(base+'/worker-health.html')
  await flow(page,'extension-csp')
  await browser.close();server.shutdown()
  print('Packaged extension page + CSP verified; installation in Microsoft Edge requires a Windows/Edge check.')

asyncio.run(main())
