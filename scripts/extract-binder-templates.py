from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
import json,base64,math
import argparse
parser=argparse.ArgumentParser(description='Extract source 7cm/9cm binder geometry and logos')
parser.add_argument('workbook',type=Path)
parser.add_argument('--output',type=Path,default=Path('src/data/binder-templates.js'))
args=parser.parse_args()
p=args.workbook
w=load_workbook(p)
profiles={}
for width,s in zip([70,90],w.worksheets):
 def col_width(c):
  d=next((d for d in s.column_dimensions.values() if (d.min or 0)<=c<=(d.max or 0)),None)
  chars=d.width if d else s.sheet_format.defaultColWidth or 9
  return math.floor(chars*7+5)*.75
 xs=[0]
 for c in range(2,12):xs.append(xs[-1]+col_width(c))
 ys={2:0}
 for r in range(2,60):ys[r+1]=ys[r]+(s.row_dimensions[r].height or s.sheet_format.defaultRowHeight)
 def box(address):
  c,r,c2,r2=range_boundaries(address)
  return {'x':round(xs[c-2],5),'y':round(ys[r],5),'width':round(xs[c2-1]-xs[c-2],5),'height':round(ys[r2+1]-ys[r],5)}
 im=s._images[12 if width==70 else 0];a=im.anchor
 start={'x':xs[a._from.col-1]+a._from.colOff/12700,'y':ys[a._from.row+1]+a._from.rowOff/12700}
 if hasattr(a,'to'):
  end={'x':xs[a.to.col-1]+a.to.colOff/12700,'y':ys[a.to.row+1]+a.to.rowOff/12700}
  logo={**start,'width':end['x']-start['x'],'height':end['y']-start['y']}
 else:logo={**start,'width':a.ext.cx/12700,'height':a.ext.cy/12700}
 total_h=logo['y']+logo['height']
 areas=({ 'site':'B2:K6','majorCode':'B8:C10','majorName':'D8:K10','middleCode':'B11:C13','middleName':'D11:K13','minor':'B15:K15','title':'E16:H42','number':'E43:H45','footer':'B49:K53'} if width==70 else {'site':'B2:K6','majorCode':'B7:C9','majorName':'D7:K9','middleCode':'B10:C12','middleName':'D10:K12','minor':'B14:K14','title':'F15:I43','number':'E44:J46','footer':'B48:K52'})
 fonts={k:s.cell(range_boundaries(addr)[1],range_boundaries(addr)[0]).font.sz or 11 for k,addr in areas.items()}
 # Map every explicit cell border; worksheet gridlines are not printed.
 segments={}
 rank={'hair':0,'thin':1,'medium':2,'double':3,'thick':4}
 merged_lookup={}
 for merged in s.merged_cells.ranges:
  c1,r1,c2,r2=range_boundaries(str(merged))
  if c1<2 or c2>11 or r1<2 or r2>(53 if width==70 else 52):continue
  for rr in range(r1,r2+1):
   for cc in range(c1,c2+1):merged_lookup[(rr,cc)]=(c1,r1,c2,r2)
 for r in range(2,54 if width==70 else 53):
  for c in range(2,12):
   merged=merged_lookup.get((r,c))
   if merged:
    c1,r1,c2,r2=merged
    if (c,r)!=(c1,r1):continue
   else:c1,r1,c2,r2=c,r,c,r
   cell=s.cell(r,c)
   edges={'left':(xs[c1-2],ys[r1],xs[c1-2],ys[r2+1]),'right':(xs[c2-1],ys[r1],xs[c2-1],ys[r2+1]),'top':(xs[c1-2],ys[r1],xs[c2-1],ys[r1]),'bottom':(xs[c1-2],ys[r2+1],xs[c2-1],ys[r2+1])}
   for side,coords in edges.items():
    style=getattr(cell.border,side).style
    if not style:continue
    key=tuple(round(v,5) for v in coords)
    old=segments.get(key)
    if not old or rank.get(style,1)>rank.get(old,1):segments[key]=style
 # Merge contiguous segments of the same style along each axis.
 groups={}
 for (x1,y1,x2,y2),style in segments.items():
  vertical=x1==x2;fixed=x1 if vertical else y1;lo=y1 if vertical else x1;hi=y2 if vertical else x2
  groups.setdefault((vertical,fixed,style),[]).append((lo,hi))
 lines=[]
 for (vertical,fixed,style),intervals in groups.items():
  lo,hi=sorted(intervals)[0]
  for a,b in sorted(intervals)[1:]+[(None,None)]:
   if a is not None and a<=hi+.001:hi=max(hi,b);continue
   lines.append({'xy':[fixed,lo,fixed,hi] if vertical else [lo,fixed,hi,fixed],'style':style})
   lo,hi=a,b
 profiles[width]={'name':s.title,'widthPt':xs[-1],'heightPt':total_h,'boxes':{k:box(v) for k,v in areas.items()},'fonts':fonts,'logo':logo,'logoData':'data:image/jpeg;base64,'+base64.b64encode(im._data()).decode(),'borders':lines,'fill':'#'+s['B2'].fill.fgColor.rgb[-6:],'fontFamily':s['B2'].font.name}
 print(width,'raw size pt',xs[-1],total_h,'class code %',profiles[width]['boxes']['majorCode']['width']/xs[-1]*100,'segments',len(lines))
args.output.write_text('// Extracted from GS건설 바인더표지 양식.xlsx (7cm/9cm), 2026-10-23.\n// Row heights, column widths, merged cells, explicit border styles and image anchors are preserved.\nexport const BINDER_TEMPLATES = '+json.dumps(profiles,ensure_ascii=False,separators=(',',':'))+';\n')
